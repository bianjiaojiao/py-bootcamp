#!/usr/bin/env python3
"""把 notes/ 下的 Markdown 笔记生成两本「电子书」：

1. notes/book/Python学习笔记.html  —— 单文件网页书（双击即开）
   · 左侧目录树：前言 / 章 / 篇（点哪读哪）
   · 顶部搜索框：输入关键词立刻定位知识点
   · 每页顶部「本页目录」：跳转到页内小节
   · 左右方向键 / 底部按钮：像翻书一样前后翻
2. notes/book/Python学习笔记.docx  —— Word 电子书（WPS/Word 打开）
   · 带自动目录 + 一级/二级/三级标题，导航窗格直接显示目录树

用法（仓库根目录）：
    python notes/tools/build_book.py            # 重新生成两本书
    python notes/tools/build_book.py --open     # 生成后自动打开网页书

Markdown 笔记是唯一内容来源：改完 md 运行一次本脚本即可。
"""
from __future__ import annotations

import argparse
import datetime
import html
import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent  # 仓库根目录
NOTES = ROOT / "notes"
TOPICS_DIR = NOTES / "topics"
DAILY_DIR = NOTES / "daily"
PLANS_DIR = NOTES / "plans"
BOOK_DIR = NOTES / "book"

CATEGORY_ORDER = ["基础结构", "文本处理", "容器", "数据处理", "进阶", "工程", "项目"]

# ---------------------------------------------------------------- Markdown 解析


def split_frontmatter(text: str) -> tuple[dict, str]:
    if not text.startswith("---"):
        return {}, text
    end = text.find("\n---", 3)
    if end == -1:
        return {}, text
    meta: dict = {}
    for line in text[3:end].splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        key, _, value = line.partition(":")
        key, value = key.strip(), value.strip()
        if value.startswith("[") and value.endswith("]"):
            meta[key] = [v.strip().strip("'\"") for v in value[1:-1].split(",") if v.strip()]
        else:
            meta[key] = value.strip("'\"")
    return meta, text[end + 5:]


def strip_comments(body: str) -> str:
    lines, out, in_comment = body.splitlines(), [], False
    for line in lines:
        s = line.strip()
        if in_comment:
            if "-->" in s:
                in_comment = False
            continue
        if s.startswith("<!--"):
            if "-->" not in s:
                in_comment = True
            continue
        out.append(line)
    return "\n".join(out)


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UL_RE = re.compile(r"^[-*]\s+")
OL_RE = re.compile(r"^\d+[.)]\s+")


def parse_blocks(body: str) -> list[dict]:
    lines = strip_comments(body).splitlines()
    blocks: list[dict] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "lang": lang, "text": "\n".join(code_lines)})
            continue
        m = HEADING_RE.match(stripped)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = re.split(r"(?<!\\)\|", lines[i].strip().strip("|"))
                rows.append([c.strip().replace("\\|", "|") for c in cells])
                i += 1
            # 第 1 行是表头，第 2 行是 |---|---| 分隔行，只删分隔行
            if len(rows) >= 2:
                del rows[1]
            blocks.append({"type": "table", "rows": rows})
            continue
        if stripped.startswith("> "):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].lstrip())
                i += 1
            blocks.append({"type": "quote", "lines": quote})
            continue
        if UL_RE.match(stripped):
            items = []
            while i < n and UL_RE.match(lines[i].strip()):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append({"type": "ul", "items": items})
            continue
        if OL_RE.match(stripped):
            items = []
            while i < n and OL_RE.match(lines[i].strip()):
                items.append(re.sub(r"^\d+[.)]\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "ol", "items": items})
            continue
        if stripped in ("---", "***", "___"):
            blocks.append({"type": "hr"})
            i += 1
            continue
        # 普通段落：连续非空行合并
        para = []
        while i < n:
            s = lines[i].strip()
            if not s or s.startswith(("```", "> ", "|")) or HEADING_RE.match(s) or UL_RE.match(s) or OL_RE.match(s) or s in ("---", "***", "___"):
                break
            para.append(s)
            i += 1
        blocks.append({"type": "p", "text": " ".join(para)})
    return blocks


# ---------------------------------------------------------------- 行内解析

INLINE_RE = re.compile(r"`([^`]*)`|\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\)")


def inline_tokens(text: str) -> list[dict]:
    tokens, pos = [], 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            tokens.append({"type": "text", "text": text[pos:m.start()]})
        if m.group(1) is not None:
            tokens.append({"type": "code", "text": m.group(1)})
        elif m.group(2) is not None:
            tokens.append({"type": "bold", "children": inline_tokens(m.group(2))})
        else:
            tokens.append({"type": "link", "text": m.group(3), "href": m.group(4)})
        pos = m.end()
    if pos < len(text):
        tokens.append({"type": "text", "text": text[pos:]})
    return tokens


def plain_text(text: str) -> str:
    """去掉 markdown 行内标记，供搜索/标题用。"""
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


# ---------------------------------------------------------------- 页面组织


def read_meta(path: Path) -> tuple[dict, list[dict], str]:
    text = path.read_text(encoding="utf-8")
    meta, body = split_frontmatter(text)
    return meta, parse_blocks(body), text


def first_h1(blocks: list[dict]) -> str:
    for b in blocks:
        if b["type"] == "h" and b["level"] == 1:
            return plain_text(b["text"])
    return ""


def collect_pages() -> tuple[list[dict], list[dict]]:
    """返回 (pages, chapters)。pages 内保存 md 相对路径、标题、章节、blocks。"""
    pages: list[dict] = []

    def add(md_path: Path, chapter: str, pid: str) -> None:
        meta, blocks, raw = read_meta(md_path)
        title = str(meta.get("title") or "").strip() or first_h1(blocks) or md_path.stem
        pages.append(
            {
                "id": pid,
                "md": md_path,
                "chapter": chapter,
                "title": title,
                "keywords": meta.get("keywords") or [],
                "blocks": blocks,
                "search": plain_text(strip_comments(raw)).replace("\n", " "),
            }
        )

    # 前言
    add(NOTES / "README.md", "前言", "readme")
    if (NOTES / "INDEX.md").exists():
        add(NOTES / "INDEX.md", "前言", "index")
    if (NOTES / "TIMELINE.md").exists():
        add(NOTES / "TIMELINE.md", "前言", "timeline")

    # 错题本（放在最前面，方便随时复习）
    if (NOTES / "错题本.md").exists():
        add(NOTES / "错题本.md", "错题本", "errorbook")

    # 按分类装主题
    topics = []
    for path in sorted(TOPICS_DIR.glob("*.md")):
        if path.name.startswith("_"):
            continue
        meta, blocks, _ = read_meta(path)
        topics.append(
            {
                "path": path,
                "category": str(meta.get("category") or "未分类"),
                "title": str(meta.get("title") or first_h1(blocks) or path.stem),
                "num": int(re.match(r"(\d+)", path.stem).group(1)) if re.match(r"(\d+)", path.stem) else 999,
            }
        )
    cats = sorted({t["category"] for t in topics}, key=lambda c: CATEGORY_ORDER.index(c) if c in CATEGORY_ORDER else 99)
    chapter_no = 0
    for cat in cats:
        chapter_no += 1
        chapter_title = f"第 {chapter_no} 章 {cat}"
        for t in sorted([x for x in topics if x["category"] == cat], key=lambda x: x["num"]):
            add(t["path"], chapter_title, f"topic-{t['num']:02d}")

    # LangChain / LangGraph 专区（框架能力，与 Python 语言知识分开）
    lc_dir = NOTES / "langchain"
    if lc_dir.exists():
        lc_files = sorted(p for p in lc_dir.glob("*.md") if p.name != "README.md")
        readme = lc_dir / "README.md"
        if readme.exists():
            lc_files = [readme] + lc_files
        if lc_files:
            chapter_no += 1
            chapter_title = f"第 {chapter_no} 章 LangChain / LangGraph"
            for i, path in enumerate(lc_files, 1):
                add(path, chapter_title, f"lc-{i:02d}")

    # 每日记录
    daily_files = sorted(DAILY_DIR.glob("*.md"), reverse=True) if DAILY_DIR.exists() else []
    if daily_files:
        chapter_no += 1
        chapter_title = f"第 {chapter_no} 章 每日记录"
        for i, path in enumerate(daily_files, 1):
            add(path, chapter_title, f"daily-{i:02d}")

    # 学习计划
    plan_files = sorted(PLANS_DIR.glob("*.md"), reverse=True) if PLANS_DIR.exists() else []
    if plan_files:
        chapter_no += 1
        chapter_title = f"第 {chapter_no} 章 学习计划"
        for i, path in enumerate(plan_files, 1):
            add(path, chapter_title, f"plan-{i:02d}")

    # 附录：原始素材
    source_files = sorted((NOTES / "source").rglob("*.md")) if (NOTES / "source").exists() else []
    if source_files:
        rows = [["原始素材文件", "说明"]]
        for path in source_files:
            rows.append([path.relative_to(NOTES).as_posix(), "截图逐字转写（已整理进上方主题笔记）"])
        append_blocks = [
            {"type": "h", "level": 1, "text": "附录 · 原始素材"},
            {"type": "p", "text": "下列文件是原始截图逐字转写归档，内容已归纳进各主题笔记，保留仅为溯源："},
            {"type": "table", "rows": rows},
        ]
        pages.append(
            {
                "id": "appendix",
                "md": NOTES / "source",
                "chapter": "附录",
                "title": "附录 · 原始素材",
                "keywords": [],
                "blocks": append_blocks,
                "search": "原始素材 截图 转写 归档",
            }
        )

    # 章节结构
    chapters: list[dict] = []
    for p in pages:
        if not chapters or chapters[-1]["title"] != p["chapter"]:
            chapters.append({"title": p["chapter"], "pages": []})
        chapters[-1]["pages"].append(p["id"])
    return pages, chapters


# ---------------------------------------------------------------- HTML 生成

CSS = """
:root{--bg:#f5f3ef;--card:#fffdf9;--text:#2b2b2b;--muted:#8a857c;--line:#e7e2d8;--accent:#b4552d;--code-bg:#f4f1ea;--sidebar:#faf8f4;--hover:#f0ebe1}
body.dark{--bg:#1b1b1f;--card:#232329;--text:#e8e4dc;--muted:#9a958c;--line:#35353c;--accent:#e08a5c;--code-bg:#2b2b33;--sidebar:#202026;--hover:#2c2c34}
*{box-sizing:border-box}
html,body{margin:0;padding:0}
body{background:var(--bg);color:var(--text);font-family:"Microsoft YaHei","PingFang SC","Segoe UI",system-ui,sans-serif;line-height:1.8}
.topbar{position:sticky;top:0;z-index:20;display:flex;gap:14px;align-items:center;padding:10px 18px;background:var(--card);border-bottom:1px solid var(--line);box-shadow:0 1px 6px rgba(0,0,0,.04)}
.brand{font-weight:700;font-size:17px;white-space:nowrap;color:var(--accent)}
#q{flex:1;max-width:520px;padding:8px 12px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--text);font-size:14px;outline:none}
#q:focus{border-color:var(--accent)}
.tbtn{border:1px solid var(--line);background:var(--bg);color:var(--text);border-radius:8px;padding:7px 10px;cursor:pointer;font-size:14px}
.tbtn:hover{background:var(--hover)}
.layout{display:flex;min-height:calc(100vh - 52px)}
#sidebar{width:292px;flex:0 0 292px;background:var(--sidebar);border-right:1px solid var(--line);overflow-y:auto;max-height:calc(100vh - 52px);position:sticky;top:52px;padding:12px 10px 40px}
.ch{margin:8px 0 4px}
.ch>button{width:100%;text-align:left;background:none;border:0;color:var(--muted);font-size:12.5px;font-weight:700;letter-spacing:.04em;padding:7px 10px;cursor:pointer;border-radius:6px}
.ch>button:hover{background:var(--hover)}
.ch>button .caret{display:inline-block;transition:transform .15s;margin-right:6px;font-size:10px}
.ch.closed>button .caret{transform:rotate(-90deg)}
.ch ul{list-style:none;margin:0;padding:0 0 4px 0}
.ch.closed ul{display:none}
.ch li a{display:block;padding:6px 10px 6px 24px;color:var(--text);text-decoration:none;font-size:14px;border-radius:6px;border-left:2px solid transparent}
.ch li a:hover{background:var(--hover)}
.ch li a.active{background:var(--hover);color:var(--accent);font-weight:600;border-left-color:var(--accent)}
#results{position:absolute;top:52px;left:0;right:0;max-width:640px;margin:0 auto;background:var(--card);border:1px solid var(--line);border-radius:0 0 10px 10px;box-shadow:0 8px 24px rgba(0,0,0,.12);display:none;max-height:60vh;overflow:auto;z-index:30}
#results .r{padding:10px 16px;border-bottom:1px solid var(--line);cursor:pointer}
#results .r:hover{background:var(--hover)}
#results .r b{color:var(--accent)}
#results .r small{color:var(--muted);display:block;margin-top:2px}
main{flex:1;min-width:0;padding:26px 18px 60px;display:flex;justify-content:center}
.page{width:100%;max-width:880px;background:var(--card);border:1px solid var(--line);border-radius:14px;padding:40px 52px 52px;box-shadow:0 2px 18px rgba(0,0,0,.05)}
.crumb{color:var(--muted);font-size:13px;margin-bottom:8px}
.crumb b{color:var(--accent)}
h1{font-size:30px;line-height:1.35;margin:.2em 0 .1em;border-bottom:2px solid var(--accent);padding-bottom:12px}
h2{font-size:22px;margin:1.6em 0 .5em;padding-left:12px;border-left:4px solid var(--accent);line-height:1.4}
h3{font-size:18px;margin:1.2em 0 .4em}
h4{font-size:16px;margin:1em 0 .3em;color:var(--muted)}
p{margin:.55em 0}
a{color:var(--accent)}
code{background:var(--code-bg);border:1px solid var(--line);border-radius:4px;padding:1px 5px;font-family:Consolas,"Cascadia Mono",monospace;font-size:.9em}
pre.code{background:var(--code-bg);border:1px solid var(--line);border-radius:10px;padding:14px 16px;overflow-x:auto;line-height:1.55;margin:.8em 0}
pre.code code{background:none;border:0;padding:0;font-size:13.5px;white-space:pre}
blockquote{margin:.8em 0;padding:8px 16px;border-left:3px solid var(--accent);background:var(--hover);border-radius:0 8px 8px 0;color:var(--text)}
blockquote p{margin:.25em 0}
table{border-collapse:collapse;width:100%;margin:.9em 0;font-size:14px}
th,td{border:1px solid var(--line);padding:7px 10px;text-align:left;vertical-align:top}
th{background:var(--hover);font-weight:700}
hr{border:0;border-top:1px solid var(--line);margin:1.4em 0}
ul,ol{padding-left:1.6em;margin:.5em 0}
li{margin:.22em 0}
.deadlink{color:var(--muted);border-bottom:1px dashed var(--muted)}
#outline{margin:14px 0 22px;padding:12px 16px;border:1px dashed var(--line);border-radius:10px;background:var(--hover);font-size:13.5px}
#outline b{font-size:13px;color:var(--muted);letter-spacing:.05em}
#outline a{display:inline-block;margin:3px 12px 3px 0;color:var(--text);text-decoration:none}
#outline a:hover{color:var(--accent);text-decoration:underline}
.pager{display:flex;justify-content:space-between;gap:12px;margin-top:34px;border-top:1px solid var(--line);padding-top:16px}
.pager button{border:1px solid var(--line);background:var(--bg);color:var(--text);padding:9px 16px;border-radius:9px;cursor:pointer;font-size:14px;max-width:48%}
.pager button:hover{background:var(--hover)}
.pager button:disabled{opacity:.35;cursor:default}
@media (max-width:900px){#sidebar{position:fixed;left:0;top:52px;bottom:0;z-index:15;transform:translateX(-100%);transition:transform .2s}#sidebar.open{transform:none}.page{padding:26px 20px 40px}h1{font-size:24px}.layout{display:block}}
.navtoggle{display:none}
@media (max-width:900px){.navtoggle{display:inline-block}}
@media print{.topbar,#sidebar,#outline,.pager,#results,.navtoggle{display:none!important}main{padding:0}.page{border:0;box-shadow:none;max-width:100%;padding:0}}
"""

JS = r"""
const BOOK = __BOOK_JSON__;
const pages = BOOK.pages;
const byId = {}; pages.forEach(p => byId[p.id] = p);
const chapters = BOOK.chapters;

function buildSidebar(){
  const nav = document.getElementById('sidebar');
  nav.innerHTML = '';
  chapters.forEach(ch => {
    const div = document.createElement('div'); div.className = 'ch';
    const btn = document.createElement('button');
    btn.innerHTML = '<span class="caret">▼</span>' + esc(ch.title);
    btn.onclick = () => { div.classList.toggle('closed'); };
    div.appendChild(btn);
    const ul = document.createElement('ul');
    ch.pages.forEach(id => {
      const p = byId[id]; if(!p) return;
      const li = document.createElement('li');
      const a = document.createElement('a');
      a.href = '#' + id; a.dataset.id = id; a.textContent = p.title;
      a.onclick = (e) => { e.preventDefault(); go(id); };
      li.appendChild(a); ul.appendChild(li);
    });
    div.appendChild(ul); nav.appendChild(div);
  });
}
function esc(s){ return String(s).replace(/[&<>"]/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c])); }
function renderBlock(b){
  if(b.type === 'h'){
    if(b.level === 1) return '';
    const tag = 'h' + Math.min(b.level + 1, 6);
    return '<' + tag + '>' + b.html + '</' + tag + '>';
  }
  if(b.type === 'p') return '<p>' + b.html + '</p>';
  if(b.type === 'code') return '<pre class="code"><code>' + esc(b.text) + '</code></pre>';
  if(b.type === 'quote') return '<blockquote>' + b.html + '</blockquote>';
  if(b.type === 'ul') return '<ul>' + b.items.map(x => '<li>' + x + '</li>').join('') + '</ul>';
  if(b.type === 'ol') return '<ol>' + b.items.map(x => '<li>' + x + '</li>').join('') + '</ol>';
  if(b.type === 'table'){
    let h = '<table><thead><tr>' + b.rows[0].map(c => '<th>' + c + '</th>').join('') + '</tr></thead><tbody>';
    b.rows.slice(1).forEach(r => { h += '<tr>' + r.map(c => '<td>' + c + '</td>').join('') + '</tr>'; });
    return h + '</tbody></table>';
  }
  if(b.type === 'hr') return '<hr>';
  return '';
}
let currentIndex = -1;
function order(){ return pages.map(p => p.id); }
function go(id, push){
  const idx = pages.findIndex(p => p.id === id);
  if(idx < 0) return;
  const p = pages[idx]; currentIndex = idx;
  const main = document.getElementById('content');
  const bodyHtml = p.blocks.map(renderBlock).join('\n');
  let outlineLinks = p.blocks.filter(b => b.type === 'h' && b.level === 2).map((b, i) =>
    '<a href="#" data-sec="' + i + '">' + b.text + '</a>').join('');
  const outline = outlineLinks ? '<div id="outline"><b>📖 本页目录</b><br>' + outlineLinks + '</div>' : '';
  main.innerHTML =
    '<div class="crumb">' + esc(p.chapter) + ' <span>›</span> <b>' + esc(p.title) + '</b></div>' +
    '<h1>' + esc(p.title) + '</h1>' + outline + bodyHtml +
    '<div class="pager">' +
    '<button id="prev"' + (idx === 0 ? ' disabled' : '') + '>← 上一页</button>' +
    '<button id="next"' + (idx === pages.length - 1 ? ' disabled' : '') + '>下一页 →</button></div>';
  // 页内小节跳转
  main.querySelectorAll('#outline a').forEach(a => {
    a.onclick = (e) => {
      e.preventDefault();
      const hs = main.querySelectorAll('h2');
      const target = hs[+a.dataset.sec];
      if(target){ target.scrollIntoView({behavior:'smooth', block:'start'}); }
    };
  });
  // 内部链接（data-go）
  main.querySelectorAll('a[data-go]').forEach(a => {
    a.onclick = (e) => { e.preventDefault(); go(a.dataset.go); };
  });
  const prev = document.getElementById('prev'), next = document.getElementById('next');
  if(prev) prev.onclick = () => { if(idx > 0) go(pages[idx - 1].id); };
  if(next) next.onclick = () => { if(idx < pages.length - 1) go(pages[idx + 1].id); };
  document.querySelectorAll('#sidebar a').forEach(a => a.classList.toggle('active', a.dataset.id === id));
  if(push !== false) history.replaceState(null, '', '#' + id);
  window.scrollTo({top: 0, behavior: 'instant' in window ? 'auto' : 'auto'});
  document.getElementById('sidebar').classList.remove('open');
  hideResults();
}
function bindSearch(){
  const q = document.getElementById('q');
  const box = document.getElementById('results');
  function run(){
    const term = q.value.trim().toLowerCase();
    if(!term){ box.style.display = 'none'; box.innerHTML = ''; return; }
    const hits = [];
    pages.forEach(p => {
      const text = (p.search || '').toLowerCase();
      const title = p.title.toLowerCase();
      let score = 0, pos = -1;
      if(title.includes(term)) score += 100;
      if((p.keywords || []).some(k => k.toLowerCase().includes(term))) score += 60;
      pos = text.indexOf(term);
      if(pos >= 0) score += 10;
      if(score > 0){
        let snippet = '';
        if(pos >= 0) snippet = (p.search || '').slice(Math.max(0, pos - 30), pos + 60);
        hits.push({p, score, snippet});
      }
    });
    hits.sort((a, b) => b.score - a.score);
    if(!hits.length){ box.innerHTML = '<div class="r" style="color:var(--muted)">没有找到「' + esc(q.value) + '」相关笔记</div>'; box.style.display = 'block'; return; }
    box.innerHTML = hits.slice(0, 12).map(h =>
      '<div class="r" data-id="' + h.p.id + '"><b>' + esc(h.p.title) + '</b>' +
      '<small>' + esc(h.p.chapter) + (h.snippet ? ' · …' + esc(h.snippet) + '…' : '') + '</small></div>'
    ).join('');
    box.style.display = 'block';
    box.querySelectorAll('.r[data-id]').forEach(el => { el.onclick = () => { go(el.dataset.id); q.value = ''; run(); }; });
  }
  q.addEventListener('input', run);
  q.addEventListener('keydown', (e) => { if(e.key === 'Enter' && box.querySelector('.r[data-id]')) box.querySelector('.r[data-id]').click(); if(e.key === 'Escape'){ q.value = ''; run(); } });
  document.addEventListener('click', (e) => { if(!box.contains(e.target) && e.target !== q) hideResults(); });
}
function hideResults(){ const box = document.getElementById('results'); box.style.display = 'none'; }
document.addEventListener('keydown', (e) => {
  if(e.target.tagName === 'INPUT') return;
  if(e.key === 'ArrowLeft' && currentIndex > 0) go(pages[currentIndex - 1].id);
  if(e.key === 'ArrowRight' && currentIndex < pages.length - 1) go(pages[currentIndex + 1].id);
  if(e.key === '/'){ e.preventDefault(); document.getElementById('q').focus(); }
});
function bindTheme(){
  const btn = document.getElementById('theme');
  const saved = localStorage.getItem('book-theme');
  if(saved === 'dark') document.body.classList.add('dark');
  btn.onclick = () => {
    document.body.classList.toggle('dark');
    localStorage.setItem('book-theme', document.body.classList.contains('dark') ? 'dark' : 'light');
  };
}
document.addEventListener('DOMContentLoaded', () => {
  buildSidebar(); bindSearch(); bindTheme();
  document.getElementById('navtoggle').onclick = () => document.getElementById('sidebar').classList.toggle('open');
  const id = location.hash.slice(1);
  go(byId[id] ? id : pages[0].id, false);
});
"""


def render_page_html(page: dict, page_map: dict, page_ids: dict) -> dict:
    """把一页的 blocks 渲染成 HTML 片段（标题除外，标题由页面壳渲染）。"""

    def linker(label: str, href: str) -> str:
        if href.startswith(("http://", "https://", "mailto:")):
            return f'<a href="{html.escape(href)}" target="_blank" rel="noopener">{html.escape(label)}</a>'
        base = page["md"].parent
        target = (base / href.split("#")[0]).resolve()
        try:
            key = target.relative_to(NOTES).as_posix()
        except ValueError:
            key = None
        if key and key in page_map:
            return f'<a href="#{page_map[key]}" data-go="{page_map[key]}">{html.escape(label)}</a>'
        if href.endswith(".md"):
            return f'<span class="deadlink" title="源文件：{html.escape(key or href)}">{html.escape(label)}</span>'
        return f'<a href="{html.escape(href)}" target="_blank" rel="noopener">{html.escape(label)}</a>'

    def inline(text: str) -> str:
        out = []
        for tok in inline_tokens(text):
            if tok["type"] == "text":
                out.append(html.escape(tok["text"]))
            elif tok["type"] == "code":
                out.append("<code>" + html.escape(tok["text"]) + "</code>")
            elif tok["type"] == "bold":
                out.append("<strong>" + inline_from_tokens(tok["children"]) + "</strong>")
            else:
                out.append(linker(tok["text"], tok["href"]))
        return "".join(out)

    def inline_from_tokens(tokens: list[dict]) -> str:
        out = []
        for tok in tokens:
            if tok["type"] == "text":
                out.append(html.escape(tok["text"]))
            elif tok["type"] == "code":
                out.append("<code>" + html.escape(tok["text"]) + "</code>")
            elif tok["type"] == "bold":
                out.append("<strong>" + inline_from_tokens(tok["children"]) + "</strong>")
            else:
                out.append(linker(tok["text"], tok["href"]))
        return "".join(out)

    blocks = []
    for b in page["blocks"]:
        if b["type"] == "h":
            blocks.append({"type": "h", "level": b["level"], "text": plain_text(b["text"]), "html": inline(b["text"])})
        elif b["type"] == "p":
            blocks.append({"type": "p", "html": inline(b["text"])})
        elif b["type"] == "code":
            blocks.append({"type": "code", "lang": b["lang"], "text": b["text"]})
        elif b["type"] == "quote":
            inner = "".join("<p>" + inline(line) + "</p>" for line in b["lines"] if line.strip())
            blocks.append({"type": "quote", "html": inner})
        elif b["type"] in ("ul", "ol"):
            items = []
            for item in b["items"]:
                if item.startswith("[ ] "):
                    items.append("☐ " + inline(item[4:]))
                elif item.startswith("[x] ") or item.startswith("[X] "):
                    items.append("☑ " + inline(item[4:]))
                else:
                    items.append(inline(item))
            blocks.append({"type": b["type"], "items": items})
        elif b["type"] == "table":
            blocks.append({"type": "table", "rows": [[inline(c) for c in row] for row in b["rows"]]})
        elif b["type"] == "hr":
            blocks.append({"type": "hr"})
    return {"type": "page", "blocks": blocks}


def build_html(pages: list[dict], chapters: list[dict]) -> Path:
    page_map = {p["md"].resolve().relative_to(NOTES).as_posix(): p["id"] for p in pages if p["md"].is_file() or str(p["md"]).endswith(".md")}
    rendered = []
    for p in pages:
        body = render_page_html(p, page_map, {})
        rendered.append(
            {
                "id": p["id"],
                "title": p["title"],
                "chapter": p["chapter"],
                "keywords": p["keywords"] if isinstance(p["keywords"], list) else [p["keywords"]],
                "search": p["search"][:4000],
                "blocks": body["blocks"],
            }
        )
    book = {"pages": rendered, "chapters": chapters}
    payload = json.dumps(book, ensure_ascii=False).replace("</", "<\\/")
    js = JS.replace("__BOOK_JSON__", payload)
    doc = (
        "<!DOCTYPE html>\n<html lang=\"zh-CN\">\n<head>\n<meta charset=\"utf-8\">\n"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        "<title>Python 学习笔记 · 电子书</title>\n<style>" + CSS + "</style>\n</head>\n<body>\n"
        "<div class=\"topbar\">"
        "<button class=\"tbtn navtoggle\" id=\"navtoggle\">☰</button>"
        "<div class=\"brand\">🐍 Python 学习笔记</div>"
        "<input id=\"q\" type=\"text\" placeholder=\"搜索知识点… 例如：排序 / get / yield （按 / 聚焦）\">"
        "<button class=\"tbtn\" id=\"theme\" title=\"切换深色/浅色\">🌓</button>"
        "</div>\n"
        "<div id=\"results\"></div>\n"
        "<div class=\"layout\">\n<nav id=\"sidebar\"></nav>\n"
        "<main><div class=\"page\" id=\"content\"></div></main>\n</div>\n"
        "<script>" + js + "</script>\n</body>\n</html>\n"
    )
    BOOK_DIR.mkdir(parents=True, exist_ok=True)
    out = BOOK_DIR / "Python学习笔记.html"
    out.write_text(doc, encoding="utf-8")
    return out


# ---------------------------------------------------------------- DOCX 生成


def build_docx(pages: list[dict], chapters: list[dict]) -> Path | None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError:
        print("[WARN] 未安装 python-docx，跳过 Word 生成（pip install python-docx）")
        return None

    MONO = "Consolas"
    YAHEI = "Microsoft YaHei"

    doc = Document()

    def set_east_asia(rpr, name):
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), name)
        rfonts.set(qn("w:hAnsi"), name)
        rfonts.set(qn("w:eastAsia"), name)

    def style_font(style, name, size, bold=False, color=None):
        style.font.name = name
        style.font.size = Pt(size)
        style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor(*color)
        set_east_asia(style.element.get_or_add_rPr(), name)

    # 全局字体
    style_font(doc.styles["Normal"], YAHEI, 10.5)
    for lvl, size in ((1, 20), (2, 16), (3, 13.5), (4, 12)):
        style_font(doc.styles[f"Heading {lvl}"], YAHEI, size, bold=True, color=(0x8B, 0x3A, 0x1A))

    # 封面
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Python 学习笔记")
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.name = YAHEI
    set_east_asia(r._r.get_or_add_rPr(), YAHEI)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("AI 专业 · Python 从零到精通 学习笔记合集")
    rs.font.size = Pt(14)
    rs.font.color.rgb = RGBColor(0x88, 0x84, 0x7C)
    r2 = doc.add_paragraph()
    r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = r2.add_run(f"生成日期：{datetime.date.today().isoformat()}   ·   手记：边角")
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(0x88, 0x84, 0x7C)
    doc.add_paragraph()
    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = tip.add_run("使用提示：WPS/Word 打开后 → 视图 → 导航窗格（大纲），即可看到一级/二级目录并点击跳转；目录页右键「更新域」可生成页码。")
    rt.font.size = Pt(9.5)
    rt.font.color.rgb = RGBColor(0x88, 0x84, 0x7C)

    # 目录（域）
    doc.add_page_break()
    doc.add_heading("目录", level=1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（右键此处 → 更新域，生成带页码的目录）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, placeholder, fld_end):
        run._r.append(el)

    # 正文
    chapter_pages: dict[str, list[dict]] = {}
    for page in pages:
        chapter_pages.setdefault(page["chapter"], []).append(page)

    def shade(par, fill="F3F5F7"):
        ppr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        ppr.append(shd)

    def add_inline_runs(par, text, size=None, color=None):
        for tok in inline_tokens(text):
            if tok["type"] == "text":
                rr = par.add_run(tok["text"])
            elif tok["type"] == "code":
                rr = par.add_run(tok["text"])
                rr.font.name = MONO
                set_east_asia(rr._r.get_or_add_rPr(), MONO)
            elif tok["type"] == "bold":
                rr = par.add_run("")
                for sub in tok["children"]:
                    if sub["type"] == "text":
                        s = par.add_run(sub["text"])
                    elif sub["type"] == "code":
                        s = par.add_run(sub["text"])
                        s.font.name = MONO
                        set_east_asia(s._r.get_or_add_rPr(), MONO)
                    else:
                        s = par.add_run(sub["text"])
                    s.font.bold = True
                    if size:
                        s.font.size = Pt(size)
                    if color:
                        s.font.color.rgb = RGBColor(*color)
                continue
            else:  # link
                rr = par.add_run(tok["text"])
                rr.font.color.rgb = RGBColor(0x8B, 0x3A, 0x1A)
            if size:
                rr.font.size = Pt(size)
            if color:
                rr.font.color.rgb = RGBColor(*color)

    for chapter in chapters:
        doc.add_page_break()
        doc.add_heading(chapter["title"], level=1)
        for pid in chapter["pages"]:
            page = next(p for p in pages if p["id"] == pid)
            doc.add_heading(page["title"], level=2)
            if page["chapter"] == "前言" and page["id"] == "readme":
                pass
            for b in page["blocks"]:
                if b["type"] == "h":
                    if b["level"] == 1:
                        continue
                    doc.add_heading(plain_text(b["text"]), level=min(b["level"] + 1, 4))
                elif b["type"] == "p":
                    par = doc.add_paragraph()
                    add_inline_runs(par, b["text"])
                elif b["type"] == "code":
                    par = doc.add_paragraph()
                    par.paragraph_format.left_indent = Pt(12)
                    par.paragraph_format.space_before = Pt(4)
                    par.paragraph_format.space_after = Pt(4)
                    par.paragraph_format.line_spacing = 1.15
                    shade(par)
                    for idx, line in enumerate(b["text"].splitlines() or [""]):
                        rr = par.add_run(line)
                        rr.font.name = MONO
                        rr.font.size = Pt(8.5)
                        set_east_asia(rr._r.get_or_add_rPr(), MONO)
                        if idx < len(b["text"].splitlines()) - 1:
                            rr.add_break()
                elif b["type"] == "quote":
                    for line in b["lines"]:
                        if not line.strip():
                            continue
                        par = doc.add_paragraph()
                        par.paragraph_format.left_indent = Pt(16)
                        par.paragraph_format.space_before = Pt(2)
                        par.paragraph_format.space_after = Pt(2)
                        add_inline_runs(par, line, color=(0x55, 0x55, 0x55))
                        for rr in par.runs:
                            rr.font.italic = True
                elif b["type"] in ("ul", "ol"):
                    style = "List Bullet" if b["type"] == "ul" else "List Number"
                    for item in b["items"]:
                        if item.startswith("[ ] "):
                            item = "☐ " + item[4:]
                        elif item.startswith("[x] "):
                            item = "☑ " + item[4:]
                        par = doc.add_paragraph(style=style)
                        add_inline_runs(par, item)
                elif b["type"] == "table":
                    if not b["rows"]:
                        continue
                    cols = max(len(r) for r in b["rows"])
                    table = doc.add_table(rows=len(b["rows"]), cols=cols)
                    table.style = "Table Grid"
                    for ri, row in enumerate(b["rows"]):
                        for ci, cell in enumerate(row):
                            if ci >= cols:
                                continue
                            pcell = table.cell(ri, ci).paragraphs[0]
                            add_inline_runs(pcell, cell)
                            if ri == 0:
                                for rr in pcell.runs:
                                    rr.font.bold = True
                    doc.add_paragraph()
                elif b["type"] == "hr":
                    doc.add_paragraph()

    BOOK_DIR.mkdir(parents=True, exist_ok=True)
    out = BOOK_DIR / "Python学习笔记.docx"
    doc.save(str(out))
    return out


# ---------------------------------------------------------------- 入口


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    parser = argparse.ArgumentParser(description="把 notes 下的 Markdown 笔记生成为 HTML/Word 电子书")
    parser.add_argument("--open", action="store_true", help="生成后自动打开网页书")
    parser.add_argument("--no-docx", action="store_true", help="只生成网页书")
    args = parser.parse_args()

    pages, chapters = collect_pages()
    if not pages:
        print("[ERROR] 没有找到任何笔记页面")
        return 1

    html_path = build_html(pages, chapters)
    print(f"[OK] 网页书：{html_path.relative_to(ROOT)}（{len(pages)} 页 / {len(chapters)} 章）")

    if not args.no_docx:
        docx_path = build_docx(pages, chapters)
        if docx_path:
            print(f"[OK] Word 书：{docx_path.relative_to(ROOT)}")

    if args.open:
        subprocess.run(["cmd", "/c", "start", "", str(html_path)], check=False)
        print("[OK] 已用默认浏览器打开网页书")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
